import streamlit as st
import pandas as pd
import json
import os

st.title("File Upload")

uploaded_file = st.file_uploader("Upload any file (CSV, Excel, JSON, TXT, etc.)", type=None)

if uploaded_file is not None:
    save_col, _ = st.columns([1, 5])
    save_result = None  # Track result for message display

    with save_col:
        if st.button("Save File to Data Folder"):
            data_dir = os.path.join(os.path.dirname(__file__), "..", "data")
            data_dir = os.path.abspath(data_dir)
            os.makedirs(data_dir, exist_ok=True)
            save_path = os.path.join(data_dir, uploaded_file.name)
            try:
                uploaded_file.seek(0)
                with open(save_path, "wb") as out_file:
                    out_file.write(uploaded_file.read())
                save_result = ("success", f"File saved to {save_path}")
            except Exception as e:
                save_result = ("error", f"Failed to save file: {e}")
    
    if save_result:
        msg_type, msg = save_result
        if msg_type == "success":
            st.success(msg)
        else:
            st.error(msg)
    file_type = uploaded_file.type
    st.write(f"**File type detected:** {file_type}")
    file_preview = None
    try:
        if file_type == "text/csv":
            df = pd.read_csv(uploaded_file)
            st.dataframe(df)
            uploaded_file.seek(0)
            file_preview = df
        elif file_type in ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel"]:
            df = pd.read_excel(uploaded_file)
            st.dataframe(df)
            uploaded_file.seek(0)
            file_preview = df
        elif file_type == "application/json":
            data = json.load(uploaded_file)
            st.json(data)
            uploaded_file.seek(0)
            file_preview = data
        elif file_type.startswith("text/"):
            content = uploaded_file.read().decode("utf-8")
            st.text_area("File Content", content, height=300)
            uploaded_file.seek(0)
            file_preview = content
        elif file_type == "application/pdf":
            try:
                import fitz  # PyMuPDF
                uploaded_file.seek(0)
                pdf_bytes = uploaded_file.read()
                doc = fitz.open(stream=pdf_bytes, filetype="pdf")
                text = ""
                for page in doc:
                    text += page.get_text()
                st.text_area("PDF Content (text only)", text, height=300)
                file_preview = text
            except ImportError:
                st.warning("PyMuPDF is not installed. Please install it to preview PDF files.")
        elif file_type in ["application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/msword"]:
            try:
                from docx import Document
                uploaded_file.seek(0)
                doc = Document(uploaded_file)
                text = "\n".join([para.text for para in doc.paragraphs])
                st.text_area("Word Document Content", text, height=300)
                file_preview = text
            except ImportError:
                st.warning("python-docx is not installed. Please install it to preview Word files.")
        elif file_type in ["application/vnd.openxmlformats-officedocument.presentationml.presentation", "application/vnd.ms-powerpoint"]:
            try:
                from pptx import Presentation
                uploaded_file.seek(0)
                prs = Presentation(uploaded_file)
                slide_texts = []
                for i, slide in enumerate(prs.slides):
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            slide_texts.append(f"Slide {i+1}: {shape.text}")
                text = "\n\n".join(slide_texts) if slide_texts else "No text found in slides."
                st.text_area("PowerPoint Content (text only)", text, height=300)
                file_preview = text
            except ImportError:
                st.warning("python-pptx is not installed. Please install it to preview PowerPoint files.")
        else:
            st.info("File uploaded, but preview is not supported for this file type.")
    except Exception as e:
        st.error(f"Error reading file: {e}")
else:
    st.info("Please upload a file to see its content.")
